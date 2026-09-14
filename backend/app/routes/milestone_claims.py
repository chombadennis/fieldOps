import os
from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy.orm import Session
from typing import List, Optional, Dict, Any
from ..db.database import get_db
from ..models.milestone_claim import MilestoneClaimDocument, MilestoneClaimItem
from ..models.project import Project
from ..models.contract import Contract
from ..models.project_integration import ProjectIntegration
from ..schemas import platform as platform_schemas
from sqlalchemy.sql import func
from pydantic import BaseModel

router = APIRouter(
    prefix="/projects/{project_id}/milestone_claims",
    tags=["Milestone Claim Documents"]
)

@router.get("")
def get_milestone_claims_documents(
    project_id: int,
    contract_id: Optional[int] = None,
    db: Session = Depends(get_db)
):
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    if contract_id is not None:
        contract = db.query(Contract).filter(Contract.id == contract_id, Contract.project_id == project_id).first()
        if not contract:
            raise HTTPException(status_code=400, detail="Contract does not belong to this project")
    else:
        general_contract = db.query(Contract).filter(Contract.project_id == project_id, Contract.contract_type == "GENERAL").first()
        if general_contract:
            contract_id = general_contract.id

    query = db.query(MilestoneClaimDocument)
    if hasattr(MilestoneClaimDocument, 'project_id'):
        query = query.filter(MilestoneClaimDocument.project_id == project_id)
        
    if hasattr(MilestoneClaimDocument, 'contract_id') and contract_id is not None:
        query = query.filter(MilestoneClaimDocument.contract_id == contract_id)

    docs = query.order_by(MilestoneClaimDocument.created_at.desc()).all()
    
    integration_ids = [d.integration_id for d in docs if d.integration_id]
    integrations_dict = {}
    if integration_ids:
        from ..models.project_integration import ProjectIntegration
        integrations = db.query(ProjectIntegration).filter(ProjectIntegration.id.in_(integration_ids)).all()
        integrations_dict = {i.id: i for i in integrations}

    result = []
    for d in docs:
        uploaded_by = None
        cloud_email = None
        if d.integration_id and d.integration_id in integrations_dict:
            integration = integrations_dict[d.integration_id]
            uploaded_by = integration.user_id
            if integration.meta_data:
                cloud_email = integration.meta_data.get('cloud_email')

        # Fetch associated line items
        items = db.query(MilestoneClaimItem).filter(MilestoneClaimItem.document_id == d.id).all()
        items_list = [
            {
                "id": item.id,
                "activity_id": item.activity_id,
                "description": item.description,
                "percentage_complete_this_period": item.percentage_complete_this_period,
                "amount_claimed_this_period": item.amount_claimed_this_period,
                "amount_certified": item.amount_certified,
                "status": item.status,
                "values_map": item.values_map,
            }
            for item in items
        ]

        result.append({
            "id": d.id,
            "project_id": getattr(d, 'project_id', project_id),
            "contract_id": getattr(d, 'contract_id', None),
            "title": d.name,
            "file_url": d.file_url or "",
            "file_type": d.file_type or "unknown",
            "department": "milestone_claims",
            "origin": d.origin or "file_upload",
            "integration_id": d.integration_id,
            "is_linked": getattr(d, 'is_linked', True),
            "linked_at": d.linked_at.isoformat() if getattr(d, 'linked_at', None) else None,
            "created_at": d.created_at.isoformat() if getattr(d, 'created_at', None) else None,
            "uploaded_by": uploaded_by,
            "cloud_email": cloud_email,
            # Milestone-specific fields
            "claim_number": d.claim_number,
            "valuation_date": d.valuation_date,
            "status": d.status,
            "payment_status": d.payment_status,
            "gross_amount_claimed": d.gross_amount_claimed or 0.0,
            "retention_deducted": d.retention_deducted or 0.0,
            "net_amount_due": d.net_amount_due or 0.0,
            "values_map": d.values_map or {},
            "items": items_list,
        })
    return result

@router.post("", response_model=platform_schemas.Document)
def create_milestone_claims_document(
    project_id: int,
    doc_in: platform_schemas.DocumentCreate,
    db: Session = Depends(get_db)
):
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    contract_id = doc_in.contract_id
    if contract_id is not None:
        contract = db.query(Contract).filter(Contract.id == contract_id, Contract.project_id == project_id).first()
        if not contract:
            raise HTTPException(status_code=400, detail="Contract does not belong to this project")
    else:
        general_contract = db.query(Contract).filter(Contract.project_id == project_id, Contract.contract_type == "GENERAL").first()
        if not general_contract:
            general_contract = Contract(project_id=project_id, name="General Contract", contract_type="GENERAL")
            db.add(general_contract)
            db.commit()
            db.refresh(general_contract)
        contract_id = general_contract.id

    existing_doc = None
    query = db.query(MilestoneClaimDocument)
    if hasattr(MilestoneClaimDocument, 'project_id'):
        query = query.filter(MilestoneClaimDocument.project_id == project_id)
    if hasattr(MilestoneClaimDocument, 'contract_id'):
        query = query.filter(MilestoneClaimDocument.contract_id == contract_id)
        
    if doc_in.cloud_file_id and hasattr(MilestoneClaimDocument, 'cloud_file_id'):
        existing_doc = query.filter(MilestoneClaimDocument.cloud_file_id == doc_in.cloud_file_id).first()
    
    if not existing_doc and doc_in.file_url:
        existing_doc = query.filter(MilestoneClaimDocument.file_url == doc_in.file_url).first()

    if existing_doc:
        existing_doc.name = doc_in.title
        existing_doc.file_url = doc_in.file_url
        existing_doc.file_type = doc_in.file_type
        if hasattr(existing_doc, 'file_size'):
            existing_doc.file_size = doc_in.file_size
        if doc_in.integration_id:
            existing_doc.integration_id = doc_in.integration_id
            
        if hasattr(existing_doc, 'is_linked'):
            existing_doc.is_linked = True
            existing_doc.linked_at = func.now()
            existing_doc.unlinked_at = None
            
        db.commit()
        db.refresh(existing_doc)
        d = existing_doc
        return platform_schemas.Document(
            id=d.id,
            project_id=getattr(d, 'project_id', project_id),
            contract_id=getattr(d, 'contract_id', None),
            title=d.name,
            file_url=d.file_url or "",
            file_type=d.file_type or "unknown",
            department="milestone_claims",
            file_size=getattr(d, 'file_size', 0),
            cloud_file_id=getattr(d, 'cloud_file_id', None),
            origin=d.origin or "file_upload",
            integration_id=d.integration_id,
            is_linked=getattr(d, 'is_linked', True),
            linked_at=getattr(d, 'linked_at', None),
            unlinked_at=getattr(d, 'unlinked_at', None),
            created_at=getattr(d, 'created_at', None)
        )

    kwargs = {
        "name": doc_in.title,
        "file_url": doc_in.file_url,
        "file_type": doc_in.file_type,
        "origin": doc_in.origin or "file_upload",
        "integration_id": doc_in.integration_id
    }
    if hasattr(MilestoneClaimDocument, 'project_id'):
        kwargs["project_id"] = project_id
    if hasattr(MilestoneClaimDocument, 'contract_id'):
        kwargs["contract_id"] = contract_id
    if hasattr(MilestoneClaimDocument, 'cloud_file_id'):
        kwargs["cloud_file_id"] = doc_in.cloud_file_id
    if hasattr(MilestoneClaimDocument, 'file_size'):
        kwargs["file_size"] = doc_in.file_size
    if hasattr(MilestoneClaimDocument, 'is_linked'):
        kwargs["is_linked"] = True
        kwargs["linked_at"] = func.now()

    new_doc = MilestoneClaimDocument(**kwargs)
    db.add(new_doc)
    db.commit()
    db.refresh(new_doc)
    d = new_doc
    return platform_schemas.Document(
        id=d.id,
        project_id=getattr(d, 'project_id', project_id),
        contract_id=getattr(d, 'contract_id', None),
        title=d.name,
        file_url=d.file_url or "",
        file_type=d.file_type or "unknown",
        department="milestone_claims",
        file_size=getattr(d, 'file_size', 0),
        cloud_file_id=getattr(d, 'cloud_file_id', None),
        origin=d.origin or "file_upload",
        integration_id=d.integration_id,
        is_linked=getattr(d, 'is_linked', True),
        linked_at=getattr(d, 'linked_at', None),
        unlinked_at=getattr(d, 'unlinked_at', None),
        created_at=getattr(d, 'created_at', None)
    )

@router.post("/{document_id}/unlink")
def unlink_milestone_claims_document(
    project_id: int,
    document_id: int,
    db: Session = Depends(get_db)
):
    doc = db.query(MilestoneClaimDocument).filter(MilestoneClaimDocument.id == document_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")
        
    if hasattr(doc, 'project_id') and doc.project_id != project_id:
        raise HTTPException(status_code=400, detail="Document does not belong to this project")

    if hasattr(doc, 'is_linked'):
        doc.is_linked = False
        doc.unlinked_at = func.now()
        db.commit()
        db.refresh(doc)
    return {"status": "unlinked"}

@router.delete("/{document_id}")
def delete_milestone_claims_document(
    project_id: int,
    document_id: int,
    db: Session = Depends(get_db)
):
    doc = db.query(MilestoneClaimDocument).filter(MilestoneClaimDocument.id == document_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")
        
    if hasattr(doc, 'project_id') and doc.project_id != project_id:
        raise HTTPException(status_code=400, detail="Document does not belong to this project")

    if getattr(doc, 'integration_id', None):
        db.query(ProjectIntegration).filter(ProjectIntegration.id == doc.integration_id).delete(synchronize_session=False)

    db.delete(doc)
    db.commit()
    return {"status": "deleted", "message": "Document data and integration permanently deleted from database"}


# --- Integration Endpoints ---

class MilestoneClaimPreviewRequest(BaseModel):
    project_id: int
    provider: str
    spreadsheet_id: str
    filename: str
    refresh_token: Optional[str] = None
    selected_sheets: Optional[List[str]] = None

class MilestoneClaimUpdateRequest(BaseModel):
    claim_number: Optional[str] = None
    valuation_date: Optional[str] = None
    status: Optional[str] = None
    payment_status: Optional[str] = None
    values_map: Optional[Dict[str, Any]] = None

class MilestoneClaimCommitRequest(BaseModel):
    project_id: int
    contract_id: Optional[int] = None
    integration_id: Optional[int] = None
    file_url: Optional[str] = None
    title: Optional[str] = "Milestone Claim"
    metrics: Dict[str, Any] = {}
    items: List[Dict[str, Any]] = []

async def _download_cloud_file(provider: str, file_id: str, refresh_token: str) -> str:
    import httpx
    import tempfile
    from ..utils.security import decrypt_token
    from ..services.integrations.google_sheets import refresh_google_access_token
    from ..services.integrations.onedrive import refresh_onedrive_access_token

    decrypted = decrypt_token(refresh_token)
    if provider == "google_sheets":
        access_token = await refresh_google_access_token(decrypted)
        headers = {"Authorization": f"Bearer {access_token}"}
        
        # Check file metadata to determine if it's a native Google Workspace file
        meta_url = f"https://www.googleapis.com/drive/v3/files/{file_id}?fields=mimeType,name"
        async with httpx.AsyncClient(timeout=30.0) as client:
            meta_resp = await client.get(meta_url, headers=headers)
            if meta_resp.status_code == 200:
                meta_data = meta_resp.json()
                mime_type = meta_data.get("mimeType", "")
                
                if mime_type == "application/vnd.google-apps.spreadsheet":
                    url = f"https://www.googleapis.com/drive/v3/files/{file_id}/export?mimeType=application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                elif mime_type == "application/vnd.google-apps.document":
                    url = f"https://www.googleapis.com/drive/v3/files/{file_id}/export?mimeType=application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                elif mime_type == "application/vnd.google-apps.presentation":
                    url = f"https://www.googleapis.com/drive/v3/files/{file_id}/export?mimeType=application/vnd.openxmlformats-officedocument.presentationml.presentation"
                else:
                    url = f"https://www.googleapis.com/drive/v3/files/{file_id}?alt=media"
            else:
                url = f"https://www.googleapis.com/drive/v3/files/{file_id}?alt=media"

    elif provider == "onedrive":
        access_token = await refresh_onedrive_access_token(decrypted)
        url = f"https://graph.microsoft.com/v1.0/me/drive/items/{file_id}/content"
        headers = {"Authorization": f"Bearer {access_token}"}
    else:
        raise HTTPException(status_code=400, detail="Unsupported cloud provider")

    fd, temp_path = tempfile.mkstemp()
    async with httpx.AsyncClient(timeout=120.0) as client:
        resp = await client.get(url, headers=headers, follow_redirects=True)
        if resp.status_code != 200:
            os.close(fd)
            os.remove(temp_path)
            error_detail = resp.text
            try:
                import json
                error_data = json.loads(resp.text)
                if "error" in error_data and "message" in error_data["error"]:
                    error_detail = error_data["error"]["message"]
            except Exception:
                pass
            raise HTTPException(status_code=resp.status_code, detail=f"Cloud download failed: {error_detail}")
        with os.fdopen(fd, 'wb') as f:
            f.write(resp.content)
            
    return temp_path

async def _extract_content_and_process(
    temp_path: str,
    filename: str,
    project_name: str,
    selected_sheets: Optional[List[str]] = None
) -> Dict[str, Any]:
    import pandas as pd
    import docx
    from ..services.integrations.milestone_ai_engine import process_milestone_claim_with_ai

    ext = os.path.splitext(filename.lower())[1]
    if not ext:
        ext = ".xlsx"
    
    try:
        if ext == ".pdf":
            return await process_milestone_claim_with_ai(
                file_path=temp_path,
                mime_type="application/pdf",
                project_name=project_name
            )
        elif ext in [".docx", ".doc"]:
            doc = docx.Document(temp_path)
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        full_text.append(" | ".join(row_cells))
            text_content = "\n".join(full_text)
            return await process_milestone_claim_with_ai(
                text_content=text_content,
                project_name=project_name
            )
        elif ext in [".xlsx", ".xls"]:
            xls = pd.ExcelFile(temp_path)
            sheets_data = {}
            
            target_sheets = selected_sheets
            if not target_sheets:
                target_sheets = xls.sheet_names[:5]
            else:
                target_sheets = [s for s in target_sheets if s in xls.sheet_names]
                if not target_sheets:
                    target_sheets = xls.sheet_names[:5]
                    
            for sheet_name in target_sheets:
                df = pd.read_excel(xls, sheet_name=sheet_name)
                rows = [df.columns.tolist()] + df.values.tolist()
                sheets_data[sheet_name] = rows
            return await process_milestone_claim_with_ai(
                sheets_data=sheets_data,
                project_name=project_name
            )
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported file format '{ext}'")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to parse document format '{ext}': {str(e)}")


@router.post("/validate")
async def validate_milestone_claim_document(
    project_id: int,
    request: MilestoneClaimPreviewRequest,
    db: Session = Depends(get_db)
):
    """
    Validates if the linked file contains a valid Milestone Claim document.
    """
    existing_link = db.query(ProjectIntegration).filter(
        ProjectIntegration.project_id == project_id,
        ProjectIntegration.spreadsheet_id == request.spreadsheet_id,
        ProjectIntegration.module != "milestone_claims_rejected"
    ).first()
    if existing_link and existing_link.module != "milestone_claims":
        raise HTTPException(
            status_code=400,
            detail=f"This document is already linked under the {existing_link.module.upper()} tab. Duplicate linkages are prohibited."
        )

    existing_rejected = db.query(ProjectIntegration).filter(
        ProjectIntegration.project_id == project_id,
        ProjectIntegration.spreadsheet_id == request.spreadsheet_id,
        ProjectIntegration.module == "milestone_claims_rejected"
    ).first()
    if existing_rejected:
        db.delete(existing_rejected)
        db.flush()

    ref_token = request.refresh_token
    if not ref_token:
        integration = db.query(ProjectIntegration).filter(
            ProjectIntegration.project_id == project_id,
            ProjectIntegration.provider == request.provider
        ).first()
        if integration:
            ref_token = integration.refresh_token

    if not ref_token:
        raise HTTPException(status_code=400, detail="OAuth credentials missing.")

    temp_path = await _download_cloud_file(request.provider, request.spreadsheet_id, ref_token)
    try:
        project = db.query(Project).filter(Project.id == project_id).first()
        proj_name = project.name if project else ""
        
        extracted_data = await _extract_content_and_process(
            temp_path,
            request.filename,
            proj_name,
            selected_sheets=request.selected_sheets
        )
        
        if not extracted_data.get("is_milestone_document"):
            doc_type = extracted_data.get("identified_document_type", "Unknown Document")
            
            db_integration = ProjectIntegration(
                project_id=project_id,
                provider=request.provider,
                spreadsheet_id=request.spreadsheet_id,
                refresh_token=ref_token or "rejected_no_token",
                sheet_name="Rejected",
                module="milestone_claims_rejected",
                meta_data={"identified_document_type": doc_type, "rejection_reason": doc_type}
            )
            db.add(db_integration)
            db.commit()

            return {"valid": False, "reason": f"Expected Milestone Claim, identified as '{doc_type}'"}
            
        if existing_rejected:
            db.commit()
            
        return {"valid": True, "reason": None, "extracted_data": extracted_data}
    finally:
        if os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except PermissionError:
                pass

@router.post("/preview")
async def preview_milestone_claim_extraction(
    project_id: int,
    request: MilestoneClaimPreviewRequest,
    db: Session = Depends(get_db)
):
    """
    Previews structured Milestone Claim extraction for human-in-the-loop validation.
    """
    ref_token = request.refresh_token
    if not ref_token:
        integration = db.query(ProjectIntegration).filter(
            ProjectIntegration.project_id == project_id,
            ProjectIntegration.provider == request.provider
        ).first()
        if integration:
            ref_token = integration.refresh_token

    if not ref_token:
        raise HTTPException(status_code=400, detail="OAuth credentials missing.")

    temp_path = await _download_cloud_file(request.provider, request.spreadsheet_id, ref_token)
    try:
        project = db.query(Project).filter(Project.id == project_id).first()
        proj_name = project.name if project else ""
        
        extracted_data = await _extract_content_and_process(
            temp_path,
            request.filename,
            proj_name,
            selected_sheets=request.selected_sheets
        )
        return extracted_data
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)

@router.post("/commit")
async def commit_milestone_claim_extraction(
    project_id: int,
    request: MilestoneClaimCommitRequest,
    db: Session = Depends(get_db)
):
    """
    Saves document details and overwrites parsed items to milestone_claim_items in Postgres.
    """
    contract_id = request.contract_id
    if not contract_id:
        gen_contract = db.query(Contract).filter(
            Contract.project_id == project_id,
            Contract.contract_type == "GENERAL"
        ).first()
        if not gen_contract:
            gen_contract = Contract(project_id=project_id, name="General Contract", contract_type="GENERAL")
            db.add(gen_contract)
            db.commit()
            db.refresh(gen_contract)
        contract_id = gen_contract.id

    doc = None
    if request.integration_id:
        doc = db.query(MilestoneClaimDocument).filter(
            MilestoneClaimDocument.project_id == project_id,
            MilestoneClaimDocument.integration_id == request.integration_id
        ).first()
    
    if not doc and request.file_url:
        doc = db.query(MilestoneClaimDocument).filter(
            MilestoneClaimDocument.project_id == project_id,
            MilestoneClaimDocument.file_url == request.file_url
        ).first()

    if not doc:
        doc = MilestoneClaimDocument(
            project_id=project_id,
            contract_id=contract_id,
            name=request.title or "Milestone Claim",
            file_url=request.file_url,
            origin="cloud_integration" if request.integration_id else "file_upload",
            integration_id=request.integration_id,
            file_type="xlsx" if (request.file_url and (".xls" in request.file_url.lower() or ".csv" in request.file_url.lower())) else "pdf",
            gross_amount_claimed=float(request.metrics.get("gross_amount_claimed", 0.0)),
            retention_deducted=float(request.metrics.get("retention_deducted", 0.0)),
            net_amount_due=float(request.metrics.get("net_amount_due", 0.0)),
            values_map=request.metrics.get("values_map", {})
        )
        db.add(doc)
        db.flush()
    else:
        doc.name = request.title or doc.name
        doc.file_url = request.file_url or doc.file_url
        doc.gross_amount_claimed = float(request.metrics.get("gross_amount_claimed", 0.0))
        doc.retention_deducted = float(request.metrics.get("retention_deducted", 0.0))
        doc.net_amount_due = float(request.metrics.get("net_amount_due", 0.0))
        doc.values_map = request.metrics.get("values_map", {})
        
        db.query(MilestoneClaimItem).filter(MilestoneClaimItem.document_id == doc.id).delete()
        db.flush()

    for item in request.items:
        db_item = MilestoneClaimItem(
            document_id=doc.id,
            activity_id=item.get("activity_id"),
            description=item.get("description", ""),
            percentage_complete_this_period=float(item.get("percentage_complete_this_period") or 0.0),
            amount_claimed_this_period=float(item.get("amount_claimed_this_period") or 0.0),
            values_map=item.get("values_map")
        )
        db.add(db_item)

    db.commit()
    db.refresh(doc)

    if request.integration_id:
        integration = db.query(ProjectIntegration).filter(ProjectIntegration.id == request.integration_id).first()
        if integration:
            integration.module = "milestone_claims"
            db.commit()

    return {
        "status": "success",
        "message": f"Successfully committed Milestone Claim '{doc.name}' with {len(request.items)} items.",
        "document_id": doc.id
    }

@router.put("/{document_id}")
def update_milestone_claim_document(
    project_id: int,
    document_id: int,
    req: MilestoneClaimUpdateRequest,
    db: Session = Depends(get_db)
):
    doc = db.query(MilestoneClaimDocument).filter(
        MilestoneClaimDocument.project_id == project_id,
        MilestoneClaimDocument.id == document_id
    ).first()

    if not doc:
        raise HTTPException(status_code=404, detail="Milestone Claim not found")

    if req.claim_number is not None:
        doc.claim_number = req.claim_number
    if req.valuation_date is not None:
        doc.valuation_date = req.valuation_date
    if req.status is not None:
        doc.status = req.status
    if req.payment_status is not None:
        doc.payment_status = req.payment_status
    if req.values_map is not None:
        doc.values_map = req.values_map

    db.commit()
    db.refresh(doc)
    
    return {"status": "success", "message": "Updated milestone claim"}
